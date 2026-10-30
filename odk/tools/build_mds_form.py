"""Build the farmer ODK form whose every choice list comes from Master Data.

    python odk/tools/build_mds_form.py --pack ../openg2p-data/packs/ETH-OAN         --geography <export of the environment's Master Data geography>

ODK Collect cannot read Master Data while a form is filled in, so the form
carries the lists at build time, exactly as Master Data serves them: the
choice names are the list's value codes (what the registry stores) and the
labels its displays. The registry's ODK transform then passes the codes
through unchanged, with no alias table between the form and the registry.

Taken from the v3 form (odk/tools/build_v3.py): every question, group, repeat,
relevance, constraint and translation. Replaced: the rows of every choice list
that is a Master Data list (the table below), and the region/zone/woreda/kebele
lists, which become the geography as external CSVs (name and code are the
level_value_id Master Data serves, which the web intake stores too; parent is
the parent's). Left as
they are: yes/no answers, consent, and the cooperative and union names, which
are not Master Data lists.

The cooperative and union pickers filter on the region the farmer chose, so their
parent key is re-keyed from the v3 region number (1, 3, 4, 8) to the region's
Master Data code (ET01, ET03, ET04, ET08). The bridge is the region's name: the
v3 form's own region list gives number -> name, and the name is matched, with
case, spacing and punctuation ignored, to the geography's regions. The cooperatives
come from odk/legacy-geo/PrimaryCoopList.legacy.csv (the v3 file, kept as it was) and
are written to odk/media/PrimaryCoopList.csv; the unions are the inline
choices_name_union rows. A row whose region has no match is left out of the published
list (a code is never invented) and is listed in the report; the build refuses
to finish if any parent key that remains is not a region of the geography.

A translation (Amharic, Afaan Oromoo) is carried over from the v3 choice that
maps to the same code, by the registry's own mapping (odk/transform.py), so the
enumerator sees the wording they know; a code the old form never offered gets
the Master Data display in every language until a translation is added.

Geography belongs to each environment's Master Data, not to the pack: an
environment migrated from an older catalogue keeps its own unit ids, while an
environment seeded from the pack serves the pack's P-codes (New Dev does since
the commons deploy of 28 October 2026: ET01 ... ET01010101, 1362 woredas,
19535 kebeles). --geography takes a directory
with that catalogue's levels.json and values.json, in the pack's format, read
from the environment's Master Data (g2p_geo_levels, g2p_geo_level_values). It
defaults to the pack's own geography, for an environment seeded from the pack.
The CSVs must be rebuilt for each environment the form is published to.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import OrderedDict
from copy import copy
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "farmer-extension/src"))
from openg2p_registry_farmer_extension.odk import transform  # noqa: E402

SOURCE = REPO / "odk/ATI_Farmers_Profile_ODK_Form_v3.xlsx"
TARGET = REPO / "odk/ATI_Farmers_Profile_ODK_Form_v4.xlsx"
MEDIA = REPO / "odk/media/mds"
COOPS_SOURCE = REPO / "odk/legacy-geo/PrimaryCoopList.legacy.csv"
COOPS_TARGET = REPO / "odk/media/PrimaryCoopList.csv"
# Rows that stay offered in every region: the filter is "region=${region} or region='<key>'".
OTHER_KEYS = {"choices_name_primary_cooperative": "other_primary_coop", "choices_name_union": "other_union"}
LABELS = ("label::English (en)", "label::Amharic (am)", "label::Afaan Oromoo (om)")
LEVELS = ("region", "zone", "woreda", "kebele")

# XLSForm choice list -> (Master Data list, how a v3 choice name becomes a code, codes left out)
MDS_LISTS: dict[str, tuple[str, callable, set[str]]] = {
    "choices_gender": ("GENDER", lambda v: transform.code(v), set()),
    "choices_member_gender": ("GENDER", lambda v: transform.code(v), set()),
    "choices_language": ("LANGUAGE", lambda v: transform.code(v, transform.LANGUAGES), set()),
    "choices_farming_type": ("FARMING_TYPE", lambda v: transform.code(v, {"CROP_FARMING": "CROP", "LIVESTOCK_FARMING": "LIVESTOCK", "MIXED_FARMING": "MIXED"}), set()),
    "choices_marital_status": ("MARITAL_STATUS", lambda v: transform.code(v), set()),
    "choices_education_level": ("EDUCATION_LEVEL", lambda v: transform.code(v, {"READ_WRITE": "CAN_READ_AND_WRITE"}), set()),
    "choices_income_source": ("SOURCE_OF_INCOME", lambda v: transform.code(v, transform.INCOME), set()),
    "choices_primary_commodity": ("CROP_COMMODITY", lambda v: transform.code(v, transform.CROPS, "CROP_"), set()),
    "choices_crops_list": ("CROP_COMMODITY", lambda v: transform.code(v, transform.CROPS, "CROP_"), set()),
    "choices_farmer_role": ("FARMER_CLUSTER_ROLE", lambda v: transform.code(v), set()),
    "choices_land_ownership": ("LAND_OWNERSHIP_TYPE", lambda v: transform.code(v), set()),
    "oan_crop_season": ("CROP_SEASON", lambda v: transform.code(v, prefix="SEASON_"), set()),
    "choices_crop_water_source": ("WATER_SOURCE", lambda v: transform.code(v, transform.WATER), set()),
    "choices_livestock_water_source": ("WATER_SOURCE", lambda v: transform.code(v, transform.WATER), set()),
    "choices_livestocks": ("LIVESTOCK_TYPE", lambda v: transform.code(v, prefix="LSTK_"), set()),
    "choices_machinery": ("FARM_MACHINERY", lambda v: transform.code(v, transform.MACHINERY), set()),
    "choices_finance": ("FINANCE_ACCESS_TYPE", lambda v: transform.code(v), set()),
    "oan_identity_document_type": ("IDENTITY_DOCUMENT_TYPE", lambda v: transform.code(v), set()),
    # A member's relationship to the head; the head is never a member's answer.
    "choices_household_relationship": ("RELATIONSHIP_TO_HEAD", lambda v: transform.code(v, transform.RELATIONSHIP), {"SELF"}),
}


def read_lists(pack: Path) -> dict[str, dict]:
    lists = {}
    for folder in (pack / "codelists", pack / "domains"):
        for path in sorted(folder.rglob("*.json")):
            data = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(data, dict) and "attribute_id" in data:
                lists[data["attribute_id"]] = data
    return lists


def read_values(pack: Path) -> dict[str, list[dict]]:
    levels = {level["level_id"]: level["level_mnemonic"] for level in json.loads((pack / "levels.json").read_text(encoding="utf-8"))}
    units: dict[str, list[dict]] = {level: [] for level in LEVELS}
    for value in json.loads((pack / "values.json").read_text(encoding="utf-8")):
        level = levels[value["level_id"]]
        if level in units:
            units[level].append(value)
    return units


def choice_rows(old_rows: list[dict], mds_list: dict, to_code, left_out: set[str]) -> tuple[list[dict], dict]:
    """The list's rows as Master Data has them, with the old translations."""
    carried = {}
    for row in old_rows:
        try:
            mapped = to_code(row["name"])
        except Exception:
            mapped = None
        if mapped:
            carried.setdefault(mapped, row)
    rows, report = [], {"codes": 0, "translated": 0, "dropped_choices": []}
    for value in sorted(mds_list["values"], key=lambda v: (v.get("sort_order") or 0, v["value_code"])):
        code = value["value_code"]
        if code in left_out:
            continue
        old = carried.get(code)
        labels = {LABELS[0]: value["value_display"]}
        for label in LABELS[1:]:
            labels[label] = old.get(label) if old and old.get(label) else value["value_display"]
        rows.append({"name": code, **labels})
        report["codes"] += 1
        report["translated"] += bool(old)
    report["dropped_choices"] = sorted(str(r["name"]) for r in old_rows if to_code(r["name"]) not in {v["value_code"] for v in mds_list["values"]})
    return rows, report


COLUMNS = ("relevant", "required", "constraint", "calculation", "trigger", "choice_filter", "default")


def literal_patterns(field: str):
    """The forms in which an expression compares a select to a choice name."""
    name = re.escape(field)
    return (re.compile(r"(\$\{" + name + r"\}\s*!?=\s*)'([^']*)'"),
            re.compile(r"(selected\(\s*\$\{" + name + r"\}\s*,\s*)'([^']*)'"),
            re.compile(r"'([^']*)'(\s*!?=\s*\$\{" + name + r"\})"))


def rewrite_literals(survey, columns: dict, renames: dict) -> list[dict]:
    """Move expressions from the old choice names to the Master Data codes.

    The v3 expressions compare a select to the old XLSForm name
    (${farming_type} = 'mixed_farming'); the list now offers the value code
    (MIXED), so a rule left as it was never fires and the questions it
    requires stay optional. renames maps each select question to
    {old name: code}; every comparison of that question to an old name is
    rewritten. Returns what changed, for the report.
    """
    changed = []
    for row in range(2, survey.max_row + 1):
        for column in COLUMNS:
            if column not in columns:
                continue
            cell = survey.cell(row, columns[column])
            if not isinstance(cell.value, str):
                continue
            text = cell.value
            for field, mapping in renames.items():
                if not mapping:
                    continue
                before, after, flipped = literal_patterns(field)
                text = before.sub(lambda m: f"{m.group(1)}'{mapping.get(m.group(2), m.group(2))}'", text)
                text = after.sub(lambda m: f"{m.group(1)}'{mapping.get(m.group(2), m.group(2))}'", text)
                text = flipped.sub(lambda m: f"'{mapping.get(m.group(1), m.group(1))}'{m.group(2)}", text)
            if text != cell.value:
                changed.append({"row": row, "name": survey.cell(row, columns["name"]).value,
                                "column": column, "was": cell.value, "now": text})
                cell.value = text
    return changed


def select_lists(survey, columns: dict) -> dict[str, str]:
    """Question name -> choice list, for the inline select_one / select_multiple questions."""
    selects = {}
    for row in range(2, survey.max_row + 1):
        kind = survey.cell(row, columns["type"]).value
        name = survey.cell(row, columns["name"]).value
        parts = kind.split() if isinstance(kind, str) else []
        if name and len(parts) > 1 and parts[0] in ("select_one", "select_multiple"):
            selects[name] = parts[1]
    return selects


def check_literals(survey, columns: dict, choice_names: dict) -> list[str]:
    """Every literal an expression compares a Master Data select to must be one of its choices."""
    selects = select_lists(survey, columns)
    problems = []
    for row in range(2, survey.max_row + 1):
        for column in COLUMNS:
            if column not in columns:
                continue
            value = survey.cell(row, columns[column]).value
            if not isinstance(value, str):
                continue
            for field, list_name in selects.items():
                if list_name not in choice_names:
                    continue
                for pattern, group in zip(literal_patterns(field), (2, 2, 1)):
                    for match in pattern.finditer(value):
                        literal = match.group(group)
                        if literal and literal not in choice_names[list_name]:
                            name = survey.cell(row, columns["name"]).value
                            problems.append(f"row {row} {name} {column}: {field} compared to '{literal}', which {list_name} does not offer")
    return problems


def normalise(name) -> str:
    """A place name with case, accents, spacing and punctuation ignored."""
    text = unicodedata.normalize("NFKD", str(name or "")).casefold()
    return re.sub(r"[^0-9a-z\u1200-\u137f]+", "", text)


def region_bridge(old_regions: list[dict], regions: list[dict]) -> tuple[dict[str, str], dict[str, str]]:
    """v3 region number -> Master Data region code, by name.

    Returns (mapping, unmatched) where unmatched maps an old key to the reason it has no code.
    """
    by_name: dict[str, list[str]] = {}
    for unit in regions:
        by_name.setdefault(normalise(unit["display_name"]), []).append(unit["level_value_id"])
    mapping, unmatched = {}, {}
    for row in old_regions:
        key, name = str(row["name"]), row.get("label::English (en)")
        found = by_name.get(normalise(name), [])
        if len(found) == 1:
            mapping[key] = found[0]
        else:
            unmatched[key] = (f"region '{name}' is not in the geography" if not found
                              else f"region '{name}' matches {len(found)} regions: {', '.join(found)}")
    return mapping, unmatched


def rekey(rows: list[dict], mapping: dict[str, str], unmatched: dict[str, str], other: str, label: str) -> tuple[list[dict], list[dict]]:
    """The rows under their region code; rows whose region has no code are left out and returned."""
    kept, left_out = [], []
    for row in rows:
        key = str(row["region"])
        if key == other:
            kept.append(row)
        elif key in mapping:
            kept.append({**row, "region": mapping[key]})
        else:
            left_out.append({"list": label, "old_key": key, "name": row["name"],
                             "reason": unmatched.get(key, f"region key {key} is not a region of the v3 form")})
    return kept, left_out


def check_parent_keys(rows: list[dict], region_codes: set[str], other: str, label: str) -> list[str]:
    """Every cooperative or union must hang under a region the geography has (or the 'other' key)."""
    return [f"{label} '{row['name']}': parent key '{row['region']}' is not a region of the geography"
            for row in rows if str(row["region"]) != other and str(row["region"]) not in region_codes]


LIVESTOCK_QUESTIONS = ("animal", "num_animals", "livestock_water_source")


def livestock_for_mixed(survey, columns: dict) -> list[dict]:
    """Mixed farmers keep animals too: the livestock questions are required for them as well.

    The v3 rules named livestock farming only (the group itself is shown for everything
    but crop farming). Applies to the farmer's questions and the household member's copies.
    """
    changed = []
    for row in range(2, survey.max_row + 1):
        name = survey.cell(row, columns["name"]).value
        for prefix, field in (("", "farming_type"), ("hh_member_", "hh_member_farming_type")):
            if name not in [prefix + q for q in LIVESTOCK_QUESTIONS]:
                continue
            for column in ("relevant", "required"):
                cell = survey.cell(row, columns[column])
                if not isinstance(cell.value, str):
                    continue
                only = f"${{{field}}} = 'LIVESTOCK'"
                both = f"(${{{field}}} = 'LIVESTOCK' or ${{{field}}} = 'MIXED')"
                text = cell.value
                if only in text and f"${{{field}}} = 'MIXED'" not in text:
                    text = text.replace(only, both)
                if text != cell.value:
                    changed.append({"row": row, "name": name, "column": column, "was": cell.value, "now": text})
                    cell.value = text
    return changed


# The registry's rules for an answer, said on the device (Collect and Enketo
# refuse the answer there). ODK Central does not run these on a submission
# posted through its API; the registry refuses it then, and the ingested form
# waits as a draft with the reason (required_fields.py). Enketo's regex is
# JavaScript's, without Unicode classes, so the scripts are code-point ranges.
ETHIOPIC = r"\u1200-\u135A\u135D-\u135F"  # Ge'ez letters and combining marks; not its punctuation or numerals
NAME_LETTERS = r"'\u2019-"  # apostrophe, right single quote, hyphen (name_rules.py); an escape, as pyxform turns a typed quote into an apostrophe
ENGLISH_LETTERS = "a-zA-Z" + NAME_LETTERS
OWN_LETTERS = "a-zA-Z" + ETHIOPIC + NAME_LETTERS  # Latin (Qubee) or Ge'ez
EN = "constraint_message::English (en)"


def _words(letters: str, count: int | None = None) -> str:
    word = f"[{letters}]+"
    return "^" + (" ".join([word] * count) if count else f"{word}( {word})*") + "$"


def _phone_rule(earlier: list[str]) -> str:
    return " and ".join([r"regex(., '^251[79][0-9]{8}$')"] + [f". != ${{{name}}}" for name in earlier])


def registry_rules() -> dict[str, dict[str, str]]:
    """Question name -> the columns it is given (name_rules.py, derived.py,
    19_phone_rules.sql, farmer_phone.py, the geography's parents)."""
    rules: dict[str, dict[str, str]] = {}
    name_message = "Enter letters only, up to 100 characters. For Tigrinya write the name in Ge'ez (Ethiopic) letters."
    phone_message = ("Enter a mobile number: 251, then 9 or 7, then 8 more digits, for example 251911223344. "
                     "Do not enter the same number twice.")
    for prefix, language in (("", "language"), ("hh_member_", "hh_member_language")):
        tigrinya = f"not(selected(${{{language}}}, 'TIGRINYA'))"
        for part in ("first", "father", "grandfather"):
            rules[f"{prefix}{part}_name_english"] = {
                "constraint": f'regex(., "{_words(ENGLISH_LETTERS)}") and string-length(.) <= 100',
                EN: "Enter English letters only (an apostrophe or hyphen is allowed), with single spaces "
                    "between words, up to 100 characters."}
            rules[f"{prefix}{part}_name_other"] = {
                "constraint": (f"({tigrinya} or regex(., '{_words(ETHIOPIC)}')) and "
                               f'regex(., "{_words(OWN_LETTERS)}") and string-length(.) <= 100'),
                EN: name_message}
            rules[f"{prefix}{part}_name_amharic"] = {
                "constraint": f"regex(., '{_words(ETHIOPIC)}') and string-length(.) <= 100",
                EN: "Enter Amharic (Ethiopic script) letters only, with single spaces between words."}
        primary, other = f"{prefix}primary_phone_number", f"{prefix}other_phone_number"
        rules[primary] = {"constraint": _phone_rule([]), EN: phone_message}
        rules[other] = {"constraint": _phone_rule([primary]), EN: phone_message}
        rules[f"{prefix}secondary_phone_number"] = {"constraint": _phone_rule([primary, other]), EN: phone_message}
    rules["other_member_name_other"] = {
        "constraint": ("(not(selected(${language}, 'TIGRINYA')) or "
                       f"regex(., '{_words(ETHIOPIC, 3)}')) and " f'regex(., "{_words(OWN_LETTERS, 3)}")'),
        EN: ("Enter three names separated by single spaces: first name, father's name, grandfather's name. "
             "For Tigrinya write them in Ge'ez (Ethiopic) letters.")}
    moved = "Select again: this is not in the place selected above."
    for level, parent in (("zone", "region"), ("woreda", "zone"), ("kebele", "woreda")):
        rules[level] = {"constraint": f"count(instance('{level}')/root/item[name=current()][parent=${{{parent}}}]) > 0",
                        EN: moved}
    # 238 woredas have no kebele in Master Data's geography.
    rules["kebele"]["required"] = "count(instance('kebele')/root/item[parent=${woreda}]) > 0"
    rules["member_reference_id"] = {
        "constraint": r"regex(., '^[A-Z]{3}-\d{4}-\d{4}-\d{4}-\d{4}-\d{6}$')",
        EN: ("Enter the Reference ID as 3 uppercase letters, four groups of 4 digits and a final group of 6 digits, "
             "separated by hyphens, for example ABC-1234-5678-9012-3456-789012.")}
    rules["land_id"] = {"constraint": "count(../../land_info_repeat[land_id = current()]) = 1",
                        EN: "Each plot needs a different Land ID."}
    return rules


def apply_registry_rules(survey, columns: dict) -> list[dict]:
    """Give the questions the registry's rules (registry_rules). A message
    column of another language that is empty gets the English text."""
    rules = registry_rules()
    seen, changed = set(), []
    others = [key for key in columns if key.startswith("constraint_message::") and key != EN]
    for row in range(2, survey.max_row + 1):
        name = survey.cell(row, columns["name"]).value
        if name not in rules:
            continue
        seen.add(name)
        for column, text in rules[name].items():
            cell = survey.cell(row, columns[column])
            if cell.value != text:
                changed.append({"row": row, "name": name, "column": column, "was": cell.value, "now": text})
                cell.value = text
        for column in others:
            cell = survey.cell(row, columns[column])
            if not cell.value and EN in rules[name]:
                cell.value = rules[name][EN]
    missing = sorted(set(rules) - seen)
    if missing:
        raise SystemExit("registry rules for questions the form does not have: " + ", ".join(missing))
    return changed


def write_media(units: dict[str, list[dict]], media: Path) -> dict[str, int]:
    media.mkdir(parents=True, exist_ok=True)
    counts = {}
    for level in LEVELS:
        path = media / f"{level}.csv"
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["name", "label", "code", "parent"])
            for unit in sorted(units[level], key=lambda u: u["level_value_id"]):
                parent = unit.get("parent_level_value_id") or ""
                writer.writerow([unit["level_value_id"], unit["display_name"], unit["level_value_id"],
                                 parent if level != "region" else ""])
        counts[level] = len(units[level])
    return counts


def build(source: Path, pack: Path, target: Path, media: Path, version: str,
          geography: Path | None = None, coops_source: Path | None = None, coops_target: Path | None = None) -> dict:
    lists = read_lists(pack)
    book = openpyxl.load_workbook(source)
    choices = book["choices"]
    header = [cell.value for cell in choices[1]]
    column = {key: index + 1 for index, key in enumerate(header) if key}
    rows = [{header[i]: cell.value for i, cell in enumerate(row) if header[i] is not None and cell.value is not None}
            for row in choices.iter_rows(min_row=2)]
    rows = [row for row in rows if row.get("list_name")]
    report = OrderedDict()
    kept: list[dict] = []
    replaced: dict[str, list[dict]] = {}
    for row in rows:
        name = row["list_name"]
        if name in MDS_LISTS or name == "choices_region":
            replaced.setdefault(name, []).append(row)
        else:
            kept.append(row)
    old_names = {name: [str(r["name"]) for r in old] for name, old in replaced.items()}
    for name, (attribute_id, to_code, left_out) in MDS_LISTS.items():
        if attribute_id not in lists:
            raise SystemExit(f"{attribute_id}: not in the pack")
        new_rows, info = choice_rows(replaced.get(name, []), lists[attribute_id], to_code, left_out)
        report[name] = {"master_data_list": attribute_id, **info}
        for new in new_rows:
            kept.append({"list_name": name, **new})
    # The region list leaves the form: the four levels are external lists.
    report["choices_region"] = {"master_data_list": "geo hierarchy (region.csv)", "codes": 0}
    geography = geography or pack
    units = read_values(geography)
    bridge, no_match = region_bridge(replaced.get("choices_region", []), units["region"])
    region_codes = {u["level_value_id"] for u in units["region"]}
    unions = [row for row in kept if row["list_name"] == "choices_name_union"]
    union_other = OTHER_KEYS["choices_name_union"]
    rekeyed_unions, left_out = rekey(unions, bridge, no_match, union_other, "union")
    kept = [row for row in kept if row["list_name"] != "choices_name_union"] + rekeyed_unions
    problems = check_parent_keys(rekeyed_unions, region_codes, union_other, "union")
    coop_rows = None
    if coops_source:
        with coops_source.open(encoding="utf-8", newline="") as handle:
            coop_rows = list(csv.DictReader(handle))
        coop_other = OTHER_KEYS["choices_name_primary_cooperative"]
        coop_rows, coops_left_out = rekey(coop_rows, bridge, no_match, coop_other, "primary cooperative")
        left_out += coops_left_out
        problems += check_parent_keys(coop_rows, region_codes, coop_other, "primary cooperative")
    if problems:
        raise SystemExit("cooperative or union parents outside the geography:\n" + "\n".join(problems))
    if coop_rows is not None and coops_target:
        coops_target.parent.mkdir(parents=True, exist_ok=True)
        with coops_target.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["list_name", "name", "label", "region"], lineterminator="\n")
            writer.writeheader()
            writer.writerows(coop_rows)
    style = [copy(choices.cell(2, c)._style) for c in range(1, choices.max_column + 1)]
    choices.delete_rows(2, choices.max_row)
    for index, row in enumerate(kept, start=2):
        for c in range(1, len(style) + 1):
            choices.cell(index, c)._style = copy(style[c - 1])
        for key, value in row.items():
            if key in column:
                choices.cell(index, column[key], value)
    survey = book["survey"]
    survey_columns = {cell.value: cell.column for cell in survey[1] if cell.value}
    for row in range(2, survey.max_row + 1):
        kind = survey.cell(row, survey_columns["type"]).value
        if isinstance(kind, str) and kind.startswith("select_one choices_region"):
            survey.cell(row, survey_columns["type"], "select_one_from_file region.csv")
    # Expressions still compare these questions to the v3 choice names.
    renames: dict[str, dict[str, str]] = {}
    for field, list_name in select_lists(survey, survey_columns).items():
        if list_name not in MDS_LISTS:
            continue
        mapping = {}
        for old in old_names.get(list_name, []):
            try:
                new = MDS_LISTS[list_name][1](old)
            except Exception:
                new = None
            if new and new != old:
                mapping[old] = new
        renames[field] = mapping
    rewritten = rewrite_literals(survey, survey_columns, renames)
    livestock = livestock_for_mixed(survey, survey_columns)
    rules = apply_registry_rules(survey, survey_columns)
    choice_names: dict[str, set[str]] = {}
    for row in kept:
        choice_names.setdefault(row["list_name"], set()).add(str(row["name"]))
    problems = check_literals(survey, survey_columns, {n: s for n, s in choice_names.items() if n in MDS_LISTS})
    if problems:
        raise SystemExit("expressions compare a select to a choice it no longer offers:\n" + "\n".join(problems))
    settings = book["settings"]
    settings_columns = {cell.value: cell.column for cell in settings[1] if cell.value}
    settings.cell(2, settings_columns["version"], version)
    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(target)
    counts = write_media(units, media)
    return {"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "pack": str(pack), "geography": str(geography), "pack_version": json.loads((pack / "manifest.json").read_text(encoding="utf-8")).get("version"),
            "output": str(target), "version": version, "rewritten_expressions": rewritten, "livestock_for_mixed": livestock, "registry_rules": rules,
            "region_bridge": bridge, "cooperatives_kept": len(coop_rows) if coop_rows is not None else None,
            "unions_kept": len(rekeyed_unions), "left_out": left_out, "media": {level: f"{media / level}.csv" for level in LEVELS},
            "units": counts, "lists": report}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--pack", required=True, type=Path, help="the country pack (packs/ETH-OAN of openg2p-data)")
    parser.add_argument("--geography", type=Path, default=None,
                        help="levels.json + values.json of the target environment's Master Data geography (default: the pack's)")
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--out", type=Path, default=TARGET)
    parser.add_argument("--media", type=Path, default=MEDIA)
    parser.add_argument("--coops-source", type=Path, default=COOPS_SOURCE, help="the v3 cooperative list, keyed by v3 region number")
    parser.add_argument("--coops-out", type=Path, default=COOPS_TARGET)
    parser.add_argument("--version", default="2026.10.31.2")
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.pack, args.out, args.media, args.version, args.geography, args.coops_source, args.coops_out), indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
