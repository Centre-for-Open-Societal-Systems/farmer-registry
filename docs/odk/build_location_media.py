#!/usr/bin/env python3
"""Build the farmer form's location lists from the shared location hierarchy.

The form picks region > zone > woreda from choice lists in the XLSForm's
choices sheet (choices_region, choices_zone, choices_woreda) and the kebele
from media/KebeleList.csv. They are generated from
docker/db-seed/seed-data/geo/geo_level_values.json, the Ethiopia hierarchy the
registry's Master Data is seeded with (region-ET04, zone-ET0408,
woreda-ET040801, kebele-ET040801101001), so every place Master Data knows is
offered, under its Master Data name and parent, and nothing else.

The codes are the P-codes as numbers, as the form has always stored them
(region 4, zone 408, woreda 40801, kebele 40801101001). farmer_transform.j2
pads the kebele back to kebele-ET040801101001 (or a woreda picked with an
"other" kebele to woreda-ET040801) and sends it as geo_lowest_level_value_id;
the farmer service fills the names from Master Data. The region number also
filters the cooperative and union lists (PrimaryCoopList.csv,
choices_name_union).

Region labels keep their Amharic and Afaan Oromoo translations; zone, woreda
and kebele names have none in the hierarchy, so all three label columns carry
the Master Data name, as before.

    python docs/odk/build_location_media.py          # rewrite the lists
    python docs/odk/build_location_media.py --check  # exit 1 if they are stale

test/test_odk_location_media.py runs the check. Publishing the result is a new
form version (README.md, "Publishing").
"""

from __future__ import annotations

import csv
import io
import json
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
REPO = HERE.parent.parent
SOURCE = REPO / "docker/db-seed/seed-data/geo/geo_level_values.json"
FORM = HERE / "ATI_Farmers_Profile_ODK_Form_v2.xlsx"
KEBELES = HERE / "media" / "KebeleList.csv"

LOCATION_LISTS = ("choices_region", "choices_zone", "choices_woreda")
LABELS = ("label::English (en)", "label::Amharic (am)", "label::Afaan Oromoo (om)")
# The form's own escape hatches, kept as they are.
OTHER_WOREDA = {"name": "other", "labels": ("Other", "Other", "Other"), "zone": "other_woreda"}
OTHER_KEBELE = ["choices_kebele", "other", "Other", "", "", "other_kebele"]


def load_hierarchy(source: pathlib.Path = SOURCE) -> dict[str, dict[str, tuple[str, str | None]]]:
    """{level: {code: (name, parent code)}}, codes as P-code digits ("0408")."""
    levels: dict[str, dict[str, tuple[str, str | None]]] = {"region": {}, "zone": {}, "woreda": {}, "kebele": {}}
    for row in json.loads(source.read_text(encoding="utf-8")):
        level = row["level_id"].removeprefix("level-")
        code = row["level_value_id"].split("-ET", 1)[1]
        parent = row.get("parent_level_value_id")
        name = (row.get("display_name") or row.get("level_value_mnemonic") or "").strip()
        levels[level][code] = (name, parent.split("-ET", 1)[1] if parent else None)
    return levels


def number(code: str) -> int:
    return int(code)


def choice_rows(levels, region_translations: dict[int, tuple[str, str]]) -> list[dict]:
    """The choices_region / choices_zone / choices_woreda rows, in sheet order."""
    rows = []
    for code, (name, _) in sorted(levels["region"].items(), key=lambda kv: kv[1][0].lower()):
        am, om = region_translations.get(number(code), (name, name))
        rows.append({"list_name": "choices_region", "name": number(code), "labels": (name, am, om)})
    for code, (name, parent) in sorted(levels["zone"].items(), key=lambda kv: (kv[1][1], kv[1][0].lower(), kv[0])):
        rows.append({"list_name": "choices_zone", "name": number(code), "labels": (name, name, name),
                     "region": number(parent)})
    for code, (name, parent) in sorted(levels["woreda"].items(), key=lambda kv: (kv[1][1], kv[1][0].lower(), kv[0])):
        rows.append({"list_name": "choices_woreda", "name": number(code), "labels": (name, name, name),
                     "region": number(levels["zone"][parent][1]), "zone": number(parent)})
    rows.append({"list_name": "choices_woreda", **OTHER_WOREDA})
    return rows


def kebele_csv(levels) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["list_name", "name", "label", "region", "zone", "woreda"])
    for code, (name, woreda) in sorted(levels["kebele"].items(), key=lambda kv: (kv[1][1], kv[1][0].lower(), kv[0])):
        zone = levels["woreda"][woreda][1]
        region = levels["zone"][zone][1]
        writer.writerow(["choices_kebele", number(code), name, number(region), number(zone), number(woreda)])
    writer.writerow(OTHER_KEBELE)
    return buffer.getvalue()


def _choices_sheet(workbook):
    sheet = workbook["choices"]
    header = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
    return sheet, header


def current_choice_rows(workbook) -> tuple[list[dict], dict[int, tuple[str, str]]]:
    sheet, header = _choices_sheet(workbook)
    col = {name: header.index(name) for name in ("list_name", "name", "region", "zone", *LABELS)}
    rows, translations = [], {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values or values[col["list_name"]] not in LOCATION_LISTS:
            continue
        row = {"list_name": values[col["list_name"]], "name": values[col["name"]],
               "labels": tuple(values[col[label]] for label in LABELS)}
        if values[col["region"]] is not None:
            row["region"] = values[col["region"]]
        if values[col["zone"]] is not None:
            row["zone"] = values[col["zone"]]
        rows.append(row)
        if row["list_name"] == "choices_region":
            translations[row["name"]] = (row["labels"][1], row["labels"][2])
    return rows, translations


def write_choice_rows(workbook, rows: list[dict]) -> None:
    """Replace the location lists, which sit together in the sheet, in place."""
    sheet, header = _choices_sheet(workbook)
    col = {name: header.index(name) + 1 for name in ("list_name", "name", "region", "zone", *LABELS)}
    first = last = None
    for index, cell in enumerate(sheet["A"][1:], start=2):
        if cell.value in LOCATION_LISTS:
            first = first or index
            last = index
    sheet.delete_rows(first, last - first + 1)
    sheet.insert_rows(first, len(rows))
    for offset, row in enumerate(rows):
        number_row = first + offset
        sheet.cell(number_row, col["list_name"], row["list_name"])
        sheet.cell(number_row, col["name"], row["name"])
        for label, value in zip(LABELS, row["labels"]):
            sheet.cell(number_row, col[label], value)
        if "region" in row:
            sheet.cell(number_row, col["region"], row["region"])
        if "zone" in row:
            sheet.cell(number_row, col["zone"], row["zone"])


def main(argv: list[str]) -> int:
    import openpyxl

    levels = load_hierarchy()
    workbook = openpyxl.load_workbook(FORM)
    current, translations = current_choice_rows(workbook)
    wanted = choice_rows(levels, translations)
    kebeles = kebele_csv(levels)
    stale = []
    if current != wanted:
        stale.append(f"{FORM.name} ({', '.join(LOCATION_LISTS)})")
    if KEBELES.read_bytes().decode("utf-8") != kebeles:
        stale.append(KEBELES.name)
    if "--check" in argv:
        if stale:
            print(f"stale, rerun docs/odk/build_location_media.py: {stale}", file=sys.stderr)
            return 1
        print("location lists match the shared hierarchy")
        return 0
    if current != wanted:
        write_choice_rows(workbook, wanted)
        workbook.save(FORM)
    KEBELES.write_bytes(kebeles.encode("utf-8"))
    print("updated: " + (", ".join(stale) if stale else "nothing, already current"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
