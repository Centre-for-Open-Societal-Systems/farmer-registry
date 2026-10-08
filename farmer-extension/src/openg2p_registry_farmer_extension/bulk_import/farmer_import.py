"""CSV/XLSX bulk import for the farmer intake form.

Pure module (stdlib + optional openpyxl for .xlsx): the column catalogue is the
single source of truth for the parser, the example files and the instructions
sheet, so they cannot drift. One row = one farmer; repeating groups (lands,
crops, livestock, household members) use numbered columns (``land_1_*``).

``rows_to_submissions`` yields the same intake JSON shape produced by
``odk/templates/farmer_transform.j2`` with ``import_source = IMPORT_FILE``.
"""

import csv
import io
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from ..register_domain.services.ethiopian_calendar import ethiopic_to_gregorian

MAX_ROWS = 1000
MAX_FILE_BYTES = 10 * 1024 * 1024
N_PHONES, N_LANDS, N_CROPS, N_LIVESTOCK, N_MEMBERS = 2, 3, 3, 3, 5

GENDERS = ["MALE", "FEMALE", "UNKNOWN"]
YES_NO = ["yes", "no"]
INCOME = ["CROP_PRODUCTION", "LIVESTOCK_PRODUCTION", "GOVERNMENT_NGO_SUPPORT", "OTHERS"]
EDU = ["ILLITERATE", "CAN_READ_AND_WRITE", "BASIC", "INTERMEDIARY", "HIGHER_EDUCATION"]
MARITAL = ["SINGLE", "MARRIED", "DIVORCED", "WIDOWED"]
DISABILITY_TYPE = ["VISION", "HEARING", "MOBILITY", "COGNITION", "SELF_CARE", "COMMUNICATION"]
DISABILITY_SEV = ["NO_DIFFICULTY", "SOME_DIFFICULTY", "A_LOT_OF_DIFFICULTY", "CANNOT_DO_AT_ALL"]
LAND_OWN = ["OWNER", "TENANT", "CROP_SHARE", "FAMILY_GIFT"]
LAND_USE = ["AGRICULTURAL", "RESIDENTIAL", "GRAZING", "FOREST"]
LS_SYSTEM = ["NOMADIC_PASTORAL", "SEMI_NOMADIC", "SEDENTARY_PASTORAL", "MIXED", "INDUSTRIAL"]
CLUSTER_ROLE = ["LEAD", "DEPUTY", "SECRETARY", "ACCOUNTANT", "MEMBER"]
END_USE = ["FOOD_HUMAN_CONSUMPTION", "FEED_ANIMALS", "BIOFUELS_NONFOOD", "OTHER"]


@dataclass(frozen=True)
class Col:
    name: str
    group: str
    description: str
    example: str = ""
    allowed: list = field(default_factory=list)
    required: bool = False


def _c(name, group, desc, ex="", allowed=None, required=False):
    return Col(name, group, desc, ex, allowed or [], required)


def _build_columns():
    cols = [
        _c("first_name", "Personal", "First name (English)", "Abebe", required=True),
        _c("middle_name", "Personal", "Middle name (English)", "Kebede"),
        _c("last_name", "Personal", "Last name (English)", "Alemu"),
        _c("first_name_amh", "Personal", "First name (Amharic)", "አበበ"),
        _c("middle_name_amh", "Personal", "Middle name (Amharic)", "ከበደ"),
        _c("last_name_amh", "Personal", "Last name (Amharic)", "አለሙ"),
        _c("first_name_om", "Personal", "First name (Afaan Oromo)", "Abebe"),
        _c("middle_name_om", "Personal", "Middle name (Afaan Oromo)", "Kebede"),
        _c("last_name_om", "Personal", "Last name (Afaan Oromo)", "Alemu"),
        _c("father_first_name", "Personal", "Father's first name", "Kebede", required=True),
        _c("father_middle_name", "Personal", "Father's middle name", "Alemu"),
        _c("father_last_name", "Personal", "Father's last name", "Tadesse"),
        _c("gender", "Personal", "Gender", "MALE", GENDERS),
        _c("birth_date", "Personal", "Date of birth (Gregorian) YYYY-MM-DD", "1985-04-12"),
        _c("birth_date_ec", "Personal", "Date of birth (Ethiopian) YYYY-MM-DD", "1977-08-04"),
        _c("estimated_age", "Personal", "Age in years if birth date unknown", "40"),
        _c("marital_status", "Socio-economic", "Marital status", "MARRIED", MARITAL),
        _c("education_level", "Socio-economic", "Education level", "BASIC", EDU),
        _c("source_of_income", "Socio-economic", "Main source of income", "CROP_PRODUCTION", INCOME),
        _c("source_of_income_other", "Socio-economic", "Describe income source when OTHERS", ""),
        _c("is_psnp_user", "Socio-economic", "Productive Safety Net Programme beneficiary", "no", YES_NO),
        _c("is_household_head", "Socio-economic", "Is household head", "yes", YES_NO),
        _c("disabled", "Socio-economic", "Has a disability", "no", YES_NO),
        _c("disability_type", "Socio-economic", "Disability type (if disabled)", "", DISABILITY_TYPE),
        _c("disability_severity", "Socio-economic", "Disability severity (if disabled)", "", DISABILITY_SEV),
        _c("language_spoken", "Socio-economic", "Primary language spoken", "ENGLISH"),
        _c("local_language", "Socio-economic", "Local/community language", "Amharic"),
        _c("national_fan", "Identification", "National FAN / UID (no spaces)", "1234567890123456"),
        _c("national_rid", "Identification", "National RID", ""),
        _c("region_name", "Location", "Region", "Oromia"),
        _c("zone_name", "Location", "Zone", "Arsi"),
        _c("woreda_name", "Location", "Woreda", "Tiyo"),
        _c("kebele_name", "Location", "Kebele", "Qila"),
        _c("latitude", "Location", "Latitude (decimal degrees)", "7.5432"),
        _c("longitude", "Location", "Longitude (decimal degrees)", "39.1234"),
    ]
    for i in range(1, N_PHONES + 1):
        cols.append(_c(f"phone_{i}", "Phone", f"Phone number {i}", "0911223344" if i == 1 else ""))
    cols += [
        _c("is_primary_cooperative_member", "Membership", "Member of a primary cooperative", "yes", YES_NO),
        _c("primary_cooperative_name", "Membership", "Primary cooperative name", "Qila Farmers Coop"),
        _c("is_cooperative_union_member", "Membership", "Member of a cooperative union", "no", YES_NO),
        _c("cooperative_union_name", "Membership", "Cooperative union name", ""),
        _c("is_farmer_cluster_member", "Membership", "Member of a farmer cluster", "no", YES_NO),
        _c("farmer_cluster_role", "Membership", "Cluster role", "", CLUSTER_ROLE),
        _c("primary_commodity", "Membership", "Cluster focal commodity", ""),
        _c("fertilizer_use", "Farm inputs", "Uses fertilizer", "yes", YES_NO),
        _c("amount_fertilizer_utilized", "Farm inputs", "Fertilizer amount (kg)", "100"),
        _c("pesticide_use", "Farm inputs", "Uses pesticide", "no", YES_NO),
        _c("amount_pesticide_utilized", "Farm inputs", "Pesticide amount", ""),
        _c("insecticide_use", "Farm inputs", "Uses insecticide", "no", YES_NO),
        _c("amount_insecticide_utilized", "Farm inputs", "Insecticide amount", ""),
        _c("improved_seed_use", "Farm inputs", "Uses improved seed", "yes", YES_NO),
        _c("amount_improved_seed_utilized", "Farm inputs", "Improved seed amount (kg)", "50"),
        _c("water_source", "Farm inputs", "Water source", "RAINFED"),
        _c("access_to_machinery", "Farm inputs", "Has access to machinery", "no", YES_NO),
        _c("access_to_finance", "Farm inputs", "Has access to finance", "no", YES_NO),
    ]
    for i in range(1, N_LANDS + 1):
        ex = i == 1
        cols += [
            _c(f"land_{i}_id", f"Land {i}", f"Parcel identifier", "LND-1" if ex else ""),
            _c(f"land_{i}_ownership", f"Land {i}", "Ownership type", "OWNER" if ex else "", LAND_OWN),
            _c(f"land_{i}_area_ha", f"Land {i}", "Total area in hectares", "1.5" if ex else ""),
            _c(f"land_{i}_kebele", f"Land {i}", "Kebele of the parcel", "Qila" if ex else ""),
            _c(f"land_{i}_current_use", f"Land {i}", "Current land use", "AGRICULTURAL" if ex else "", LAND_USE),
        ]
    for i in range(1, N_CROPS + 1):
        ex = i == 1
        cols += [
            _c(f"crop_{i}_name", f"Crop {i}", "Crop / commodity", "WHEAT" if ex else ""),
            _c(f"crop_{i}_planted_date", f"Crop {i}", "Planted date YYYY-MM-DD", "2026-06-15" if ex else ""),
            _c(f"crop_{i}_season", f"Crop {i}", "Season code from configured master data", "SUMMER" if ex else ""),
            _c(f"crop_{i}_end_use", f"Crop {i}", "End use", "FOOD_HUMAN_CONSUMPTION" if ex else "", END_USE),
        ]
    for i in range(1, N_LIVESTOCK + 1):
        ex = i == 1
        cols += [
            _c(f"livestock_{i}_type", f"Livestock {i}", "Animal type", "CATTLE" if ex else ""),
            _c(f"livestock_{i}_head_count", f"Livestock {i}", "Number of animals", "4" if ex else ""),
            _c(f"livestock_{i}_system", f"Livestock {i}", "Livestock system", "MIXED" if ex else "", LS_SYSTEM),
        ]
    for i in range(1, N_MEMBERS + 1):
        ex = i == 1
        cols += [
            _c(f"member_{i}_name", f"Household member {i}", "Full name (first middle last)", "Almaz Tadesse" if ex else ""),
            _c(f"member_{i}_gender", f"Household member {i}", "Gender", "FEMALE" if ex else "", GENDERS),
            _c(f"member_{i}_birth_date", f"Household member {i}", "Birth date YYYY-MM-DD", "1988-02-01" if ex else ""),
            _c(f"member_{i}_is_disabled", f"Household member {i}", "Has a disability", "no" if ex else "", YES_NO),
        ]
    cols += [
        _c("enumerator_name", "Enumerator", "Data collector name", "Staff User"),
        _c("enumerator_user_id", "Enumerator", "Data collector id", "ENUM-001"),
        _c("data_collection_date", "Enumerator", "Collection date YYYY-MM-DD", "2026-10-01"),
    ]
    return cols


COLUMNS = _build_columns()
HEADERS = [c.name for c in COLUMNS]
_TRUE = {"yes", "y", "true", "1"}
_FALSE = {"no", "n", "false", "0"}


class RowError(ValueError):
    pass


class ParsedRow(dict):
    """A row with its original worksheet/CSV record number."""

    def __init__(self, values, number):
        super().__init__(values)
        self.number = number


def _s(row, key):
    v = row.get(key)
    return "" if v is None else str(v).strip()


def _bool(row, key):
    v = _s(row, key).lower()
    if not v:
        return None
    if v in _TRUE:
        return True
    if v in _FALSE:
        return False
    raise RowError(f"{key}: '{v}' is not yes/no")


def _num(row, key, kind=float):
    v = _s(row, key)
    if not v:
        return None
    try:
        number = float(v)
        if not math.isfinite(number) or number < 0:
            raise ValueError()
        if kind is int and not number.is_integer():
            raise ValueError()
        return kind(number)
    except (ValueError, OverflowError):
        raise RowError(f"{key}: '{v}' must be a finite non-negative {'whole number' if kind is int else 'number'}")


def _enum(row, key, allowed, default=""):
    v = _s(row, key).upper().replace(" ", "_")
    if not v:
        return default
    if v not in allowed:
        raise RowError(f"{key}: '{v}' not in {', '.join(allowed)}")
    return v


def _date(row, key):
    raw = row.get(key)
    if isinstance(raw, (date, datetime)):
        if key.endswith('_ec'):
            raise RowError(f"{key}: enter an Ethiopian date as YYYY-MM-DD text")
        return raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat()
    v = _s(row, key)
    if v:
        try:
            if not re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}", v):
                raise ValueError()
            if key.endswith('_ec'):
                ethiopic_to_gregorian(*map(int, v.split('-')))
            else:
                date.fromisoformat(v)
        except (ValueError, OverflowError):
            raise RowError(f"{key}: '{v}' must be a valid {'Ethiopian' if key.endswith('_ec') else 'Gregorian'} date (YYYY-MM-DD)")
    return v


def row_to_submission(row):
    """Map one parsed row (dict header->value) to the intake JSON."""
    if not _s(row, "first_name"):
        raise RowError("first_name is required")
    if not _s(row, "father_first_name"):
        raise RowError("father_first_name is required")

    phones = [
        {"phone_number": _s(row, f"phone_{i}"), "phone_type": "PRIMARY" if i == 1 else "SECONDARY",
         "is_primary": i == 1}
        for i in range(1, N_PHONES + 1) if _s(row, f"phone_{i}")
    ]
    reg_ids = []
    for col, id_type in (("national_fan", "UID"), ("national_rid", "RID")):
        v = _s(row, col).replace(" ", "")
        if v:
            reg_ids.append({"id_type": id_type, "value": v, "status": "VALID"})

    lands, crops, livestock, members = [], [], [], []
    for i in range(1, N_LANDS + 1):
        if any(_s(row, f"land_{i}_{k}") for k in ("id", "ownership", "area_ha", "kebele", "current_use")):
            lands.append({
                "land_id": _s(row, f"land_{i}_id") or f"LND-{i}",
                "land_ownership_type": _enum(row, f"land_{i}_ownership", LAND_OWN, "OWNER"),
                "area_in_hectare": _num(row, f"land_{i}_area_ha") or 0.0,
                "land_kebele": _s(row, f"land_{i}_kebele") or _s(row, "kebele_name"),
                "current_land_use": _enum(row, f"land_{i}_current_use", LAND_USE),
            })
    for i in range(1, N_CROPS + 1):
        if _s(row, f"crop_{i}_name"):
            crops.append({
                "commodity": _s(row, f"crop_{i}_name").upper(),
                "planted_date": _date(row, f"crop_{i}_planted_date"),
                "season": _s(row, f"crop_{i}_season").upper(),
                "end_use": _enum(row, f"crop_{i}_end_use", END_USE, "FOOD_HUMAN_CONSUMPTION"),
            })
    for i in range(1, N_LIVESTOCK + 1):
        if _s(row, f"livestock_{i}_type"):
            livestock.append({
                "livestock_type": _s(row, f"livestock_{i}_type").upper(),
                "head_count": _num(row, f"livestock_{i}_head_count", int) if _s(row, f"livestock_{i}_head_count") else 1,
                "livestock_system": _enum(row, f"livestock_{i}_system", LS_SYSTEM, "MIXED"),
            })
    for i in range(1, N_MEMBERS + 1):
        name = _s(row, f"member_{i}_name")
        if name:
            p = name.split()
            members.append({
                "first_name": p[0], "middle_name": p[1] if len(p) > 1 else "",
                "last_name": " ".join(p[2:]),
                "gender": _enum(row, f"member_{i}_gender", GENDERS, "UNKNOWN"),
                "birth_date": _date(row, f"member_{i}_birth_date"),
                "is_disabled": _bool(row, f"member_{i}_is_disabled") or False,
            })

    farm_in = {}
    for k in ("fertilizer", "pesticide", "insecticide", "improved_seed"):
        farm_in[f"{k}_use"] = _bool(row, f"{k}_use") or False
        farm_in[f"amount_{k}_utilized"] = _num(row, f"amount_{k}_utilized") or 0.0
    farm_in.update(water_source=_s(row, "water_source"),
                   access_to_machinery=_bool(row, "access_to_machinery") or False,
                   access_to_finance=_bool(row, "access_to_finance") or False)

    is_coop = _bool(row, "is_primary_cooperative_member") or False
    is_union = _bool(row, "is_cooperative_union_member") or False
    is_cluster = _bool(row, "is_farmer_cluster_member") or False
    disabled = _bool(row, "disabled") or False

    out = {
        "fr_farmer_personal_info": [{
            "first_name": _s(row, "first_name"), "middle_name": _s(row, "middle_name"),
            "last_name": _s(row, "last_name"),
            "first_name_amh": _s(row, "first_name_amh"), "middle_name_amh": _s(row, "middle_name_amh"),
            "last_name_amh": _s(row, "last_name_amh"),
            "first_name_om": _s(row, "first_name_om"), "middle_name_om": _s(row, "middle_name_om"),
            "last_name_om": _s(row, "last_name_om"),
            "father_first_name": _s(row, "father_first_name"),
            "father_middle_name": _s(row, "father_middle_name"),
            "father_last_name": _s(row, "father_last_name"),
            "phone_numbers": phones,
            "gender": _enum(row, "gender", GENDERS, "UNKNOWN"),
            "marital_status": _enum(row, "marital_status", MARITAL),
            "education_level": _enum(row, "education_level", EDU),
            "language_spoken": _s(row, "language_spoken"),
            "local_language": _s(row, "local_language"),
            "state": "PENDING",
            "import_source": "IMPORT_FILE",
        }],
        "fr_farmer_phone_numbers": phones,
        "fr_farmer_birth_information": [{
            "birth_date": _date(row, "birth_date"), "birth_date_ec": _date(row, "birth_date_ec"),
            "estimated_age": _num(row, "estimated_age", int),
        }],
        "fr_farmer_location": [{
            "region_name": _s(row, "region_name"), "zone_name": _s(row, "zone_name"),
            "woreda_name": _s(row, "woreda_name"), "kebele_name": _s(row, "kebele_name"),
            "latitude": _s(row, "latitude") or None, "longitude": _s(row, "longitude") or None,
        }],
        "fr_farmer_reg_ids": reg_ids,
        "fr_farmer_household_lookup": [{"is_household_head": _bool(row, "is_household_head") or False}],
        "fr_farmer_socio_and_health": [{
            "source_of_income": _enum(row, "source_of_income", INCOME),
            "source_of_income_other": _s(row, "source_of_income_other"),
            "is_psnp_user": _bool(row, "is_psnp_user") or False,
            "disabled": disabled,
            "disability_type": _enum(row, "disability_type", DISABILITY_TYPE) if disabled else "",
            "disability_severity": _enum(row, "disability_severity", DISABILITY_SEV) if disabled else "",
        }],
        "fr_farmer_membership": [{
            "is_primary_cooperative_member": is_coop,
            "primary_cooperative_name": _s(row, "primary_cooperative_name") if is_coop else "",
            "is_cooperative_union_member": is_union,
            "cooperative_union_name": _s(row, "cooperative_union_name") if is_union else "",
            "is_farmer_cluster_member": is_cluster,
            "farmer_cluster_role": _enum(row, "farmer_cluster_role", CLUSTER_ROLE, "MEMBER") if is_cluster else "",
            "primary_commodity": _s(row, "primary_commodity").upper() if is_cluster else "",
        }],
        "fr_household_members": members,
        "fr_farmer_enumerator": [{
            "enumerator_name": _s(row, "enumerator_name"),
            "enumerator_user_id": _s(row, "enumerator_user_id"),
            "data_collection_date": _date(row, "data_collection_date"),
        }],
    }
    if lands:
        out["intake_fr_farmer_land"] = lands
    if crops:
        out["intake_fr_farmer_crops"] = crops
    if livestock:
        out["fr_farmer_livestocks"] = livestock
    if any(_s(row, k) for k in farm_in):
        out["fr_farmer_farm_input"] = [farm_in]
    return out


def read_rows(filename, content):
    """Parse CSV/XLSX bytes into a list of header->value dicts (1 per data row)."""
    if len(content) > MAX_FILE_BYTES:
        raise RowError("File exceeds the 10 MB limit")
    name = (filename or "").lower()
    workbook = None
    try:
        if name.endswith(".csv"):
            values = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True)
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook
            # Check expanded size before openpyxl loads shared strings/styles.
            from zipfile import ZipFile
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise RowError("Expanded workbook exceeds the 50 MB limit")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.worksheets[0]
            values = sheet.iter_rows(values_only=True)
        else:
            raise RowError("Unsupported file type; upload a .csv or .xlsx file")
        header = [str(h).strip() if h is not None else "" for h in next(values, [])]
        # Excel can retain formatting in otherwise empty trailing columns.
        while header and not header[-1]:
            header.pop()
        if not header or any(not h for h in header):
            raise RowError("A non-empty header is required for every column")
        if len(header) != len(set(header)):
            raise RowError("Duplicate column headers are not allowed")
        unknown = [h for h in header if h not in HEADERS]
        if unknown:
            raise RowError("Unknown column(s): " + ", ".join(unknown))
        missing = [c.name for c in COLUMNS if c.required and c.name not in header]
        if missing:
            raise RowError("Missing required column(s): " + ", ".join(missing))
        rows = []
        for number, cells in enumerate(values, 2):
            if not any(v is not None and str(v).strip() for v in cells):
                continue
            if any(v is not None and str(v).strip() for v in cells[len(header):]):
                raise RowError(f"Row {number}: more values than column headers")
            rows.append(ParsedRow(dict(zip(header, cells)), number))
            if len(rows) > MAX_ROWS:
                raise RowError(f"Too many rows; the limit is {MAX_ROWS}")
        if not rows:
            raise RowError("File contains no farmer rows")
        return rows
    except (UnicodeDecodeError, csv.Error, BadZipFile, KeyError, IndexError, ParseError, OSError) as error:
        raise RowError("Unreadable file; use a UTF-8 CSV or valid XLSX workbook") from error
    finally:
        if workbook is not None:
            workbook.close()


def rows_to_submissions(rows):
    """Return [{row, ok, submission|errors}] without aborting on a bad row."""
    results = []
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        n = getattr(row, "number", n)
        try:
            results.append({"row": n, "ok": True, "submission": row_to_submission(row)})
        except RowError as e:
            results.append({"row": n, "ok": False, "errors": [str(e)]})
    return results


def example_rows():
    return [{c.name: c.example for c in COLUMNS}]


def build_csv():
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=HEADERS, lineterminator="\n")
    w.writeheader()
    w.writerows(example_rows())
    return ("﻿" + buf.getvalue()).encode("utf-8")


def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(HEADERS)
    ws.append([c.example for c in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ins = wb.create_sheet("Instructions")
    ins.append(["Column", "Group", "Required", "Description", "Allowed values", "Example"])
    for cell in ins[1]:
        cell.font = Font(bold=True)
    for c in COLUMNS:
        ins.append([c.name, c.group, "Yes" if c.required else "", c.description,
                    ", ".join(c.allowed), c.example])
    ins.append([])
    ins.append(["Notes: one row per farmer; delete the example row before uploading; "
                "leave unused columns blank; dates are YYYY-MM-DD; yes/no for flags; "
                f"maximum {MAX_ROWS} rows per file."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
