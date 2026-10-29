"""The farmer ODK form's choice lists are Master Data's (odk/tools/build_mds_form.py)."""
import csv
import importlib.util
import json
from pathlib import Path

import openpyxl
import pytest

REPO = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("build_mds_form", REPO / "odk/tools/build_mds_form.py")
build_mds_form = importlib.util.module_from_spec(spec)
spec.loader.exec_module(build_mds_form)

LABELS = build_mds_form.LABELS


def mds_list(attribute_id, codes):
    return {"attribute_id": attribute_id, "attribute_code": attribute_id, "attribute_display": attribute_id,
            "is_hierarchical": False, "values": [
                {"value_id": f"{attribute_id}_{c}", "value_code": c, "value_display": c.replace("_", " ").title(),
                 "sort_order": i + 1, "display_i18n": {}} for i, c in enumerate(codes)]}


@pytest.fixture
def pack(tmp_path):
    """A pack with every list the form needs, two of each geography level."""
    root = tmp_path / "ETH-TEST"
    (root / "codelists").mkdir(parents=True)
    (root / "domains/agriculture").mkdir(parents=True)
    lists = {
        "GENDER": ["FEMALE", "MALE", "OTHER"], "LANGUAGE": ["AMHARIC", "AFAAN_OROMO", "SOMALI", "TIGRINYA", "AFAR", "GEDEO"],
        "FARMING_TYPE": ["CROP", "LIVESTOCK", "MIXED", "AQUACULTURE"], "MARITAL_STATUS": ["SINGLE", "MARRIED"],
        "EDUCATION_LEVEL": ["ILLITERATE", "CAN_READ_AND_WRITE", "TERTIARY"],
        "SOURCE_OF_INCOME": ["SOI_CROP_PRODUCTION", "SOI_LIVESTOCK_PRODUCTION", "SOI_OTHERS"],
        "CROP_COMMODITY": ["CROP_TEFF", "CROP_MAIZE", "CROP_BREAD_WHEAT", "CROP_MALT_BARLEY"],
        "FARMER_CLUSTER_ROLE": ["LEAD", "MEMBER"], "LAND_OWNERSHIP_TYPE": ["OWNER", "TENANT", "HYBRID"],
        "CROP_SEASON": ["SEASON_MEHER", "SEASON_BELG"], "WATER_SOURCE": ["WS_RAINFED", "WS_WELL_GROUND", "WS_SURFACE_WATER"],
        "LIVESTOCK_TYPE": ["LSTK_CATTLE", "LSTK_GOAT"], "FARM_MACHINERY": ["MACH_TRACTOR", "MACH_SPRAYER"],
        "FINANCE_ACCESS_TYPE": ["LOAN", "SAVING"], "IDENTITY_DOCUMENT_TYPE": ["FAYDA", "PASSPORT"],
        "RELATIONSHIP_TO_HEAD": ["SELF", "SPOUSE", "CHILD", "OTHER_RELATIVE"],
    }
    for attribute_id, codes in lists.items():
        folder = "domains/agriculture" if attribute_id in ("CROP_COMMODITY", "CROP_SEASON", "FARM_MACHINERY", "FINANCE_ACCESS_TYPE") else "codelists"
        (root / folder / f"{attribute_id.lower()}.json").write_text(json.dumps(mds_list(attribute_id, codes)), encoding="utf-8")
    (root / "levels.json").write_text(json.dumps([
        {"level_id": "l0", "level_mnemonic": "country", "parent_level_id": None},
        {"level_id": "l1", "level_mnemonic": "region", "parent_level_id": "l0"},
        {"level_id": "l2", "level_mnemonic": "zone", "parent_level_id": "l1"},
        {"level_id": "l3", "level_mnemonic": "woreda", "parent_level_id": "l2"},
        {"level_id": "l4", "level_mnemonic": "kebele", "parent_level_id": "l3"}]), encoding="utf-8")
    units = [{"level_value_id": "ET", "level_id": "l0", "display_name": "Ethiopia", "parent_level_value_id": None}]
    for r, region in (("01", "Tigray"), ("03", "Amhara")):
        units.append({"level_value_id": f"ET{r}", "level_id": "l1", "display_name": region, "parent_level_value_id": "ET"})
        for z in ("01", "02"):
            units.append({"level_value_id": f"ET{r}{z}", "level_id": "l2", "display_name": f"Zone {r}{z}", "parent_level_value_id": f"ET{r}"})
            units.append({"level_value_id": f"ET{r}{z}01", "level_id": "l3", "display_name": f"Woreda {r}{z}01", "parent_level_value_id": f"ET{r}{z}"})
            units.append({"level_value_id": f"ET{r}{z}0101", "level_id": "l4", "display_name": f"Kebele {r}{z}0101", "parent_level_value_id": f"ET{r}{z}01"})
    (root / "values.json").write_text(json.dumps(units), encoding="utf-8")
    (root / "manifest.json").write_text(json.dumps({"version": "test"}), encoding="utf-8")
    return root


def test_every_master_data_list_replaces_its_choice_list(pack, tmp_path):
    out = tmp_path / "form.xlsx"
    media = tmp_path / "media"
    report = build_mds_form.build(build_mds_form.SOURCE, pack, out, media, "test-version")
    book = openpyxl.load_workbook(out, read_only=True)
    rows = list(book["choices"].iter_rows(values_only=True))
    header = list(rows[0])
    by_list = {}
    for row in rows[1:]:
        if row[0]:
            by_list.setdefault(row[0], []).append(dict(zip(header, row)))
    lists = build_mds_form.read_lists(pack)
    for name, (attribute_id, _, left_out) in build_mds_form.MDS_LISTS.items():
        expected = [v["value_code"] for v in lists[attribute_id]["values"] if v["value_code"] not in left_out]
        assert [r["name"] for r in by_list[name]] == expected, name
        for r in by_list[name]:
            assert all(r.get(label) for label in LABELS), (name, r["name"])
    # A code the old form knew keeps its translations; the member's relationship never offers SELF.
    gender = {r["name"]: r for r in by_list["choices_gender"]}
    assert gender["MALE"]["label::Amharic (am)"] != gender["OTHER"]["label::Amharic (am)"]
    assert "SELF" not in [r["name"] for r in by_list["choices_household_relationship"]]
    # Lists that are not Master Data stay as they were: yes/no, consent, unions.
    assert [r["name"] for r in by_list["choices_phone_number"]] == ["yes", "no"]
    # Unions hang under the Master Data code of their region; this pack has Tigray and Amhara only.
    assert {r["region"] for r in by_list["choices_name_union"]} == {"ET01", "ET03", "other_union"}
    assert len(by_list["choices_name_union"]) < 164
    assert "choices_region" not in by_list
    assert report["lists"]["choices_machinery"]["master_data_list"] == "FARM_MACHINERY"
    assert book["settings"].cell(2, 3).value == "test-version"
    # The geography is the pack's hierarchy, one external list per level, P-codes as name and code.
    for level, count in (("region", 2), ("zone", 4), ("woreda", 4), ("kebele", 4)):
        with (media / f"{level}.csv").open(encoding="utf-8", newline="") as handle:
            units = list(csv.DictReader(handle))
        assert len(units) == count == report["units"][level]
        assert all(u["name"] == u["code"] and u["label"] for u in units)
        if level != "region":
            assert all(u["parent"] and u["code"].startswith(u["parent"]) for u in units)
    survey = list(book["survey"].iter_rows(values_only=True))
    types = {row[1]: row[0] for row in survey[1:] if row[1]}
    for level in build_mds_form.LEVELS:
        assert types[level] == f"select_one_from_file {level}.csv"
        assert types["hh_member_" + level] == f"select_one_from_file {level}.csv"


def test_choice_rows_report_old_choices_master_data_does_not_have():
    old = [{"name": "tractor", "label::English (en)": "Tractor", "label::Amharic (am)": "ትራክተር"},
           {"name": "rice_polisher", "label::English (en)": "Rice polisher"}]
    rows, info = build_mds_form.choice_rows(old, mds_list("FARM_MACHINERY", ["MACH_TRACTOR", "MACH_SPRAYER"]),
                                            build_mds_form.MDS_LISTS["choices_machinery"][1], set())
    assert [r["name"] for r in rows] == ["MACH_TRACTOR", "MACH_SPRAYER"]
    assert rows[0]["label::Amharic (am)"] == "ትራክተር" and rows[1]["label::Amharic (am)"] == "Mach Sprayer"
    assert info == {"codes": 2, "translated": 1, "dropped_choices": ["rice_polisher"]}


def test_geography_comes_from_the_environments_master_data(pack, tmp_path):
    """An environment migrated from an older catalogue keeps its own unit ids;
    the form's location lists carry those, not the pack's."""
    geo = tmp_path / "env-geo"
    geo.mkdir()
    (geo / "levels.json").write_text(json.dumps([
        {"level_id": "level-region", "level_mnemonic": "region", "parent_level_id": None},
        {"level_id": "level-zone", "level_mnemonic": "zone", "parent_level_id": "level-region"},
        {"level_id": "level-woreda", "level_mnemonic": "woreda", "parent_level_id": "level-zone"},
        {"level_id": "level-kebele", "level_mnemonic": "kebele", "parent_level_id": "level-woreda"}]), encoding="utf-8")
    (geo / "values.json").write_text(json.dumps([
        {"level_value_id": "region-ET01", "level_id": "level-region", "display_name": "Tigray", "parent_level_value_id": None},
        {"level_value_id": "zone-ET0101", "level_id": "level-zone", "display_name": "North Western", "parent_level_value_id": "region-ET01"},
        {"level_value_id": "woreda-ET010101", "level_id": "level-woreda", "display_name": "Tahtay Adiyabo", "parent_level_value_id": "zone-ET0101"},
        {"level_value_id": "kebele-ET010101101001", "level_id": "level-kebele", "display_name": "Kebele 1", "parent_level_value_id": "woreda-ET010101"}]), encoding="utf-8")
    media = tmp_path / "media"
    report = build_mds_form.build(build_mds_form.SOURCE, pack, tmp_path / "form.xlsx", media, "v", geography=geo)
    assert report["units"] == {"region": 1, "zone": 1, "woreda": 1, "kebele": 1}
    with (media / "kebele.csv").open(encoding="utf-8", newline="") as handle:
        assert list(csv.DictReader(handle)) == [{"name": "kebele-ET010101101001", "label": "Kebele 1",
                                                 "code": "kebele-ET010101101001", "parent": "woreda-ET010101"}]
    with (media / "region.csv").open(encoding="utf-8", newline="") as handle:
        assert [r["name"] for r in csv.DictReader(handle)] == ["region-ET01"]


def survey_columns(book):
    header = [cell.value for cell in book["survey"][1]]
    return {key: index + 1 for index, key in enumerate(header) if key}


def choice_names(book):
    names = {}
    for row in book["choices"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            names.setdefault(row[0], set()).add(str(row[1]))
    return names


def test_rules_compare_selects_to_the_codes_the_lists_offer(pack, tmp_path):
    """The v3 rules named the old choices (farming_type = 'mixed_farming'); with the lists
    carrying Master Data codes they must say MIXED, or the questions they require stay optional."""
    out = tmp_path / "form.xlsx"
    report = build_mds_form.build(build_mds_form.SOURCE, pack, out, tmp_path / "media", "v")
    book = openpyxl.load_workbook(out)
    columns = survey_columns(book)
    rows = {book["survey"].cell(r, columns["name"]).value: r for r in range(2, book["survey"].max_row + 1)}
    required = book["survey"].cell(rows["primary_cooperative"], columns["required"]).value
    assert required == "${farming_type} = 'MIXED' or ${farming_type} = 'CROP'"
    assert book["survey"].cell(rows["hh_member_crop_information"], columns["relevant"]).value == "${hh_member_farming_type} != 'LIVESTOCK'"
    assert "'AFAAN_OROMO'" in book["survey"].cell(rows["first_name_other"], columns["relevant"]).value
    assert book["survey"].cell(rows["other_income_source"], columns["relevant"]).value == "selected(${income_source}, 'SOI_OTHERS')"
    assert report["rewritten_expressions"]
    lists = {n: s for n, s in choice_names(book).items() if n in build_mds_form.MDS_LISTS}
    assert build_mds_form.check_literals(book["survey"], columns, lists) == []


def test_a_rule_naming_a_choice_the_list_lacks_is_reported():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["type", "name", "relevant"])
    sheet.append(["select_one choices_farming_type", "farming_type", None])
    sheet.append(["text", "coop", "${farming_type} = 'mixed_farming'"])
    columns = {"type": 1, "name": 2, "relevant": 3}
    problems = build_mds_form.check_literals(sheet, columns, {"choices_farming_type": {"CROP", "MIXED"}})
    assert len(problems) == 1 and "mixed_farming" in problems[0]


def test_the_committed_v4_form_has_no_stale_rule():
    book = openpyxl.load_workbook(REPO / "odk/ATI_Farmers_Profile_ODK_Form_v4.xlsx")
    lists = {n: s for n, s in choice_names(book).items() if n in build_mds_form.MDS_LISTS}
    assert build_mds_form.check_literals(book["survey"], survey_columns(book), lists) == []
    settings = {c.value: book["settings"].cell(2, c.column).value for c in book["settings"][1] if c.value}
    assert settings["form_id"] == "farmer_profile" and settings["version"] == "2026.10.31.1"


COOPS = [
    {"list_name": "choices_name_primary_cooperative", "name": "Wegahita", "label": "Wegahita", "region": "1"},
    {"list_name": "choices_name_primary_cooperative", "name": "Gozamin", "label": "Gozamin", "region": "3"},
    {"list_name": "choices_name_primary_cooperative", "name": "Hiwet", "label": "Hiwet", "region": "4"},
    {"list_name": "choices_name_primary_cooperative", "name": "other", "label": "Other", "region": "other_primary_coop"},
]


def write_coops(path, rows=COOPS):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["list_name", "name", "label", "region"])
        writer.writeheader()
        writer.writerows(rows)


def test_cooperatives_and_unions_are_keyed_by_region_code(pack, tmp_path):
    """The region question stores ET01; a list still keyed 1 would leave a Tigray farmer with only Other."""
    source, target = tmp_path / "legacy.csv", tmp_path / "out" / "PrimaryCoopList.csv"
    write_coops(source)
    report = build_mds_form.build(build_mds_form.SOURCE, pack, tmp_path / "form.xlsx", tmp_path / "media", "v",
                                  coops_source=source, coops_target=target)
    with target.open(encoding="utf-8", newline="") as handle:
        published = list(csv.DictReader(handle))
    assert [(r["name"], r["region"]) for r in published] == [("Wegahita", "ET01"), ("Gozamin", "ET03"), ("other", "other_primary_coop")]
    # Oromia is not in this pack: its cooperative is left out, not given an invented code, and reported.
    left_out = [r for r in report["left_out"] if r["list"] == "primary cooperative"]
    assert [(r["old_key"], r["name"]) for r in left_out] == [("4", "Hiwet")]
    assert "Oromia" in left_out[0]["reason"]
    assert report["region_bridge"]["1"] == "ET01" and report["region_bridge"]["3"] == "ET03"
    book = openpyxl.load_workbook(tmp_path / "form.xlsx", read_only=True)
    rows = list(book["choices"].iter_rows(values_only=True))
    region = rows[0].index("region")
    unions = [r for r in rows[1:] if r[0] == "choices_name_union"]
    assert unions and {r[region] for r in unions} == {"ET01", "ET03", "other_union"}


def test_the_matching_ignores_case_spacing_and_punctuation():
    assert build_mds_form.normalise("Benishangul-Gumuz") == build_mds_form.normalise("  BENISHANGUL GUMUZ ")
    old = [{"name": 2, "label::English (en)": "South West Ethiopia"}, {"name": 9, "label::English (en)": "Nowhere"}]
    regions = [{"level_value_id": "ET11", "display_name": "South-West  ethiopia"}]
    mapping, unmatched = build_mds_form.region_bridge(old, regions)
    assert mapping == {"2": "ET11"} and "Nowhere" in unmatched["9"]


def test_a_parent_key_outside_the_geography_is_reported():
    rows = [{"name": "A", "region": "ET01"}, {"name": "B", "region": "7"}, {"name": "other", "region": "other_union"}]
    problems = build_mds_form.check_parent_keys(rows, {"ET01"}, "other_union", "union")
    assert len(problems) == 1 and "'B'" in problems[0] and "'7'" in problems[0]
    assert build_mds_form.check_parent_keys(rows[:1] + rows[2:], {"ET01"}, "other_union", "union") == []


def test_the_build_stops_when_a_cooperative_parent_is_outside_the_geography(pack, tmp_path, monkeypatch):
    """If re-keying ever let a stale key through, the guard fails the build instead of publishing it."""
    source = tmp_path / "legacy.csv"
    write_coops(source)
    monkeypatch.setattr(build_mds_form, "rekey", lambda rows, *args: (rows, []))
    with pytest.raises(SystemExit) as stopped:
        build_mds_form.build(build_mds_form.SOURCE, pack, tmp_path / "form.xlsx", tmp_path / "media", "v",
                             coops_source=source, coops_target=tmp_path / "coops.csv")
    assert "outside the geography" in str(stopped.value) and "Wegahita" in str(stopped.value)


def test_unmatched_rows_are_left_out_never_given_a_code():
    kept, left_out = build_mds_form.rekey([{"name": "A", "region": "1"}, {"name": "B", "region": "9"}],
                                          {"1": "ET01"}, {"9": "region 'Nowhere' is not in the geography"}, "other_union", "union")
    assert kept == [{"name": "A", "region": "ET01"}]
    assert left_out == [{"list": "union", "old_key": "9", "name": "B", "reason": "region 'Nowhere' is not in the geography"}]


LIVESTOCK = ("animal", "num_animals", "livestock_water_source")


def test_livestock_questions_are_required_for_mixed_farmers_too(pack, tmp_path):
    out = tmp_path / "form.xlsx"
    report = build_mds_form.build(build_mds_form.SOURCE, pack, out, tmp_path / "media", "v")
    book = openpyxl.load_workbook(out)
    columns = survey_columns(book)
    survey = book["survey"]
    rows = {survey.cell(r, columns["name"]).value: r for r in range(2, survey.max_row + 1)}
    for prefix, field in (("", "farming_type"), ("hh_member_", "hh_member_farming_type")):
        for question in LIVESTOCK:
            required = survey.cell(rows[prefix + question], columns["required"]).value
            assert f"${{{field}}} = 'LIVESTOCK'" in required and f"${{{field}}} = 'MIXED'" in required, (prefix, question)
        # The group they sit in is shown for every farmer who is not crop-only, so mixed farmers see it.
        assert survey.cell(rows[prefix + "livestock_info"], columns["relevant"]).value == f"${{{field}}} != 'CROP'"
    assert {c["name"] for c in report["livestock_for_mixed"]} == {p + q for p in ("", "hh_member_") for q in LIVESTOCK}


def test_the_livestock_rule_is_not_applied_twice():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["name", "relevant", "required"])
    sheet.append(["animal", None, "(${farming_type} = 'LIVESTOCK' or ${farming_type} = 'MIXED')"])
    assert build_mds_form.livestock_for_mixed(sheet, {"name": 1, "relevant": 2, "required": 3}) == []


def test_the_committed_cooperative_and_union_lists_use_region_codes():
    regions = {row["code"] for row in csv.DictReader((REPO / "odk/media/mds/region.csv").open(encoding="utf-8"))}
    coops = list(csv.DictReader((REPO / "odk/media/PrimaryCoopList.csv").open(encoding="utf-8")))
    assert coops and {r["region"] for r in coops} - {"other_primary_coop"} <= regions
    book = openpyxl.load_workbook(REPO / "odk/ATI_Farmers_Profile_ODK_Form_v4.xlsx", read_only=True)
    rows = list(book["choices"].iter_rows(values_only=True))
    region = rows[0].index("region")
    unions = {r[region] for r in rows[1:] if r[0] == "choices_name_union"}
    assert unions and unions - {"other_union"} <= regions
    survey = list(book["survey"].iter_rows(values_only=True))
    column = survey[0].index("required")
    required = {r[1]: r[column] for r in survey[1:] if r[1]}
    assert "'MIXED'" in required["animal"] and "'MIXED'" in required["hh_member_livestock_water_source"]
